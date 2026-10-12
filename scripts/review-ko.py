"""Produce an English/Korean workbook and CSVs from the catalogue actually shipped."""
from pathlib import Path
import json, re, csv, html
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

catalog_source = Path('js/i18n-catalog.js').read_text(encoding='utf-8')
catalog = json.loads(catalog_source.split('window.PBC_PHRASES = ',1)[1].strip().removesuffix(';'))
audit = json.loads(Path('translation-audit.json').read_text(encoding='utf-8'))
locations = {}
for row in audit:
    source = html.unescape(row['source'])
    locations.setdefault(source,set()).update(row['files'])
authored = {line.split(' ⟐ ')[0] for line in Path('locales/site.txt').read_text(encoding='utf-8').splitlines() if line and not line.startswith('#')}
overrides_path = Path('locales/ko.json')
overrides = json.loads(overrides_path.read_text(encoding='utf-8')) if overrides_path.exists() else {}
word_path = Path('locales/ko-words.json')
words = json.loads(word_path.read_text(encoding='utf-8')) if word_path.exists() else {}
headers = ['ID','English','Korean / 한국어','Used in','Translation status','Reviewer correction / 수정','Reviewer notes / 의견']
rows = []
for i, (source, values) in enumerate(sorted(catalog.items(), key=lambda entry:entry[0].lower()),1):
    korean = values.get('ko','')
    if not korean: raise ValueError('Missing Korean: '+source)
    status = 'Authored translation' if source in authored or source in overrides else 'Original identifier / proper name' if korean == source else 'Draft — native review requested'
    rows.append([f'P{i:04}', source, korean, ', '.join(sorted(locations.get(source,{'Shared UI'}))), status, '', ''])
word_rows = []
for word in sorted(set(re.findall(r"[A-Za-z]+(?:['’][A-Za-z]+)?", ' '.join(catalog))),key=str.lower):
    exact = next((v['ko'] for s,v in catalog.items() if s.lower()==word.lower()), None)
    word_rows.append([word, exact or words.get(word.lower(),''), '', ''])
if any(not row[1] for row in word_rows): raise ValueError('Word glossary has missing Korean entries')

folder = Path('docs'); folder.mkdir(exist_ok=True)
for name, columns, data in [('korean-phrase-review.csv',headers,rows),('korean-word-review.csv',['English word','Korean draft / 한국어','Reviewer correction / 수정','Reviewer notes / 의견'],word_rows)]:
    with (folder/name).open('w',encoding='utf-8-sig',newline='') as handle:
        writer = csv.writer(handle); writer.writerow(columns); writer.writerows(data)

workbook = Workbook()
notes = workbook.active; notes.title = 'Read me'
for row in [
    ['English–Korean site translation review / 영어–한국어 번역 검토'],
    ['The Phrases sheet contains the Korean text used by the website, paired with its English source.'],
    ['Use Reviewer correction and Reviewer notes to suggest changes. No approval is implied by this workbook.'],
    ['Longer new translations are drafts. The user-supplied Korean ban/status wording is preserved.'],
    ['Words is a separate vocabulary index. A word’s standalone meaning can differ from its meaning in a full phrase.'],
    ['Placeholders such as {0} represent player names, numbers, dates, or other values inserted while using the site.'],
    ['Proper names, player names, account IDs, URLs, weapon model numbers, currency codes, and literal keyboard keys can remain unchanged.'],
    ['Internal identifiers and asset filenames are not user-facing translation phrases.'],
    ['Every phrase row comes from the shipped catalogue; the same build regenerates this review.'],
    ['Phrase count',len(rows)],['Word count',len(word_rows)]
]: notes.append(row)
notes.column_dimensions['A'].width = 110; notes.column_dimensions['B'].width = 20
for cell in notes['A']: cell.alignment=Alignment(wrap_text=True,vertical='top')
for i in range(1,10): notes.row_dimensions[i].height=38

def sheet(name, columns, data, widths, table_name):
    ws=workbook.create_sheet(name); ws.append(columns)
    for row in data: ws.append(row)
    ws.freeze_panes='B2'
    for i,width in enumerate(widths,1): ws.column_dimensions[chr(64+i)].width=width
    for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='17365D')
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.alignment=Alignment(wrap_text=True,vertical='top')
        # Force text so a phrase beginning with punctuation never becomes an Excel formula.
        for cell in row:
            if isinstance(cell.value,str): cell.data_type='s'
        longest=max(len(str(cell.value or '')) for cell in row[:4])
        ws.row_dimensions[row[0].row].height=min(180,max(32,16*(longest//60+1)))
    table=Table(displayName=table_name,ref=f'A1:{chr(64+len(columns))}{len(data)+1}')
    table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
    ws.add_table(table)
    ws.sheet_view.zoomScale=85
    return ws
sheet('Phrases',headers,rows,[12,65,65,32,30,65,45],'PhrasesReview')
sheet('Words',['English word','Korean draft / 한국어','Reviewer correction / 수정','Reviewer notes / 의견'],word_rows,[35,45,45,45],'WordsReview')
workbook.save(folder/'korean-translation-review.xlsx')
print(f'Review workbook and CSVs: {len(rows)} phrases, {len(word_rows)} words.')
